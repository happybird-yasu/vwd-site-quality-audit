"""Additional source adapters. Original public sources only, no production writes."""
from __future__ import annotations
import csv
import io
import json
from pathlib import Path
import re
import zipfile
from .official_night import fetch, load, number, compact, PREFECTURES, workbook_rows
from .run import SITE


def unique(rows,key):
    out={}
    for row in rows:
        k=key(row)
        if k in out:raise ValueError('Duplicate original-source key '+str(k))
        out[k]=row
    return out


def csv_rows(raw):
    for enc in ['utf-8-sig','cp932']:
        try:return list(csv.reader(io.StringIO(raw.decode(enc))))
        except UnicodeDecodeError:pass
    raise ValueError('Unknown CSV encoding')


def literal(q,key,field,actual,original):
    r=q.result;r['expected']+=1
    if actual is None or original is None:
        r['unverified']+=1;status='unverified'
    else:
        r['compared']+=1;status='matched' if actual==original else 'mismatch'
        r['matched' if status=='matched' else 'mismatches']+=1
    if status!='matched':r['details'].append(dict(key=key,field=field,published=actual,official=original,status=status))


def mineral_index(raw):
    z=zipfile.ZipFile(io.BytesIO(raw));names=[n for n in z.namelist() if n.endswith('/mine_jp.csv')]
    if len(names)!=1:raise ValueError('Expected one official mine_jp.csv')
    rows=csv_rows(z.read(names[0]));header=rows[0]
    required={'id','mineral_j','mine_name','work_status','lat','lon','map_pub_year','map_uri'}
    if not required.issubset(header):raise ValueError('Mineral original schema changed')
    return unique([dict(zip(header,r)) for r in rows[1:] if r],lambda r:r['id'])


def minerals(c,q):
    from .checks import json_values
    source='https://www.gsj.jp/data/interim-report/GSJ_DOC_INR_073_2017_DATA.zip'
    index=mineral_index(fetch(c,source))
    q.result['scopes'].append(dict(source=source,publication=2017,
        fields=['id','longitude','latitude','work_status','map_publication_year','mineral_j','mine_name','map_uri'],
        scope='every published national mineral-map record',coordinate_precision='five decimals as published'))
    raw=fetch(c,SITE+'/maps/minerals/data/national.js','published_numeric_data')
    objects=json_values(raw.decode('utf-8'),is_json=False)[0]
    data=next(o for _,o in objects if isinstance(o,dict) and 'records' in o and 'sources' in o)
    seen=set()
    for row in data['records']:
        if len(row)!=9:raise ValueError('Unexpected published mineral record width')
        key=str(row[0])
        if key in seen:raise ValueError('Duplicate published mineral record '+key)
        seen.add(key);s=index.get(key,{});source_info=data['sources'][row[7]]
        for field,actual,original in [('longitude',row[1],s.get('lon')),('latitude',row[2],s.get('lat'))]:
            n=number(original);q.cell(key,field,actual,round(n,5) if n is not None else None)
        q.cell(key,'work_status',row[8],s.get('work_status'))
        q.cell(key,'map_publication_year',source_info.get('year'),s.get('map_pub_year'))
        for field,a,b in [('mineral_j',row[5],s.get('mineral_j')),('mine_name',row[6],s.get('mine_name')),
                          ('map_uri',source_info.get('uri'),s.get('map_uri'))]:
            literal(q,key,field,a,b.strip() if isinstance(b,str) else b)
    q.result['limitations'] += ['This checks historical source records, not present-day mineral reserves, mining operation or production.',
        'Prefecture boundary assignment, category definitions and source rows excluded from the map are outside this record-value comparison.',
        'Coordinates are verified at the published five-decimal precision, not as a fresh field survey.']


def prefecture(v):
    s=compact(v)
    for p in PREFECTURES:
        if s==compact(p) or (p!='北海道' and s==compact(p[:-1])):return p
    return '全国' if s=='全国' else None


def medical_table(raw,year):
    sheets=workbook_rows(raw);candidates=[]
    for rows in sheets:
        head=''.join(compact(v) for row in rows[:8] for v in row)
        # Older .xls omits the table title in cell values, so also require the
        # complete prefectural row coverage rather than relying on a sheet index.
        names={prefecture(v) for row in rows for v in row if prefecture(v)}
        if '歯科診療所' in head and len(names)==48:candidates.append(rows)
    if len(candidates)!=1:raise ValueError('Expected one prefectural facility-count table')
    rows=candidates[0];out={}
    for row in rows:
        matches=[(j,prefecture(v)) for j,v in enumerate(row) if prefecture(v)]
        if not matches:continue
        j,name=matches[0]
        start=next((k for k in range(j+1,len(row)) if number(row[k]) is not None),None)
        if start is None:continue
        if start+4>=len(row):raise ValueError('Medical count row truncated')
        if name in out:raise ValueError('Duplicate medical prefecture '+name)
        out[name]={'hospital':row[start],'clinic':row[start+2],'dental':row[start+4]}
    if len(out)!=48:raise ValueError(f'Expected national +47 prefectures, got{len(out)}')
    # The workbook itself must identify its own December observation date.
    text=''.join(compact(v) for rs in sheets for row in rs[:10] for v in row)
    era='令和元年' if year==2019 else '令和'+str(year-2018)+'年'
    if '12月' not in text or era not in text:raise ValueError('Medical observation year/month changed')
    return out


def medical(c,q):
    url=SITE+'/maps/medical/japan-medical-map/national-data/comparison/medical-stats.json'
    data=load(c,url)
    original={}
    for year in [2019,2021,2024]:
        yy=str(year)[2:];ext='xls' if year==2019 else 'xlsx'
        src=f'https://www.mhlw.go.jp/toukei/saikin/hw/iryosd/m{yy}/xls/is{yy}12.{ext}'
        q.result['scopes'].append(dict(source=src,year=year,as_of='December month-end provisional',
            fields=['hospital','clinic','dental'],scope='national +47 prefecture historical facility counts',unit='facilities'))
        try:original[str(year)]=medical_table(fetch(c,src),year)
        except Exception as e:original[str(year)]={};q.error(src,e)
    regions=[('全国',data['national'])]+[(p['name'],p['values']) for p in data['prefectures'].values()]
    for name,values in regions:
        for year in ['2019','2021','2024']:
            actual=values[year];actual=actual.get('values') if isinstance(actual,dict) else actual
            src=original[year].get(name,{})
            for i,field in enumerate(data['schema']):q.cell([name,year],field,actual[i],src.get(field))
    q.result['limitations'] += ['Historical counts use December month-end provisional reports. October1 annual facility survey is a different reference date and must not be substituted.',
        'Population denominators and derived per-capita rates are outside this adapter.',
        'The2025-12 individual facility points, beds, specialties, coordinates and multilingual attributes are outside this historical-count comparison.']


def port_table(raw):
    out={}
    for rows in workbook_rows(raw):
        positions=set()
        for row in rows:
            positions.update(j for j,v in enumerate(row) if compact(v)=='種別' and j>=4)
        for j in positions:
            pref=None;name=None;record=None
            for row in rows:
                if len(row)<=j+1:continue
                p=prefecture(row[j-4])
                if p and p!='全国':pref=p
                kind=compact(row[j])
                rawname=compact(row[j-2])
                # Some ports print only an inland or foreign row, without a
                # preceding total row. A new name ALWAYS starts a new record.
                if kind in ['計','外','内'] and rawname:
                    name=rawname if rawname and rawname!='合計' and '計' not in rawname else None
                    if name and pref:
                        key=(pref,name)
                        if key in out:raise ValueError('Duplicate port reference key '+str(key))
                        record={};out[key]=record
                    else:record=None
                if kind in ['計','外','内'] and record is not None:
                    record[{'計':'total','外':'foreign','内':'domestic'}[kind]]=row[j+1]
    return out


def ports(c,q):
    sources=json.loads(Path(__file__).with_name('port-sources.json').read_text())
    index={str(y):{} for y in [2019,2021,2024]}
    for s in sources:
        q.result['scopes'].append(dict(source=s['url'],year=s['year'],fields=['total','domestic','foreign'],
            unit='passengers/year',key='prefecture +port name',scope='original annual passenger tables, A/B ports'))
        try:
            parsed=port_table(fetch(c,s['url']));target=index[str(s['year'])]
            for key,val in parsed.items():
                if key in target and target[key]!=val:raise ValueError('Conflicting official port key '+str(key))
                target[key]=val
        except Exception as e:q.error(s['url'],e)
    data=load(c,SITE+'/maps/ports/passengers/ports.json')
    for row in data:
        key=(PREFECTURES[int(row['pref_code'])-1],compact(row['name_ja']))
        for year in ['2019','2021','2024']:
            src=index[year].get(key,{})
            for field in ['total','domestic','foreign']:
                q.cell([*key,year],field,row['years'].get(year,{}).get(field),src.get(field))
    q.result['limitations'] += ['Non-numeric official cells, including printed dashes, stay unverified rather than inferred as zero.',
        'Port names are matched within prefectures; unrecognized names and missing tables cannot pass.',
        'Coordinates, translated labels, excluded ports and ranking displays remain outside this passenger-count comparison.']


def lodging_facilities(raw,year):
    rows=csv_rows(raw);text=''.join(''.join(r) for r in rows[:5])
    era='平成30年度' if year==2018 else '令和元年度' if year==2019 else '令和'+str(year-2018)+'年度'
    if str(year)+'年度' not in compact(text) and era not in compact(text):raise ValueError('Lodging fiscal year changed')
    if '旅館・ホテル' not in text:raise ValueError('Wrong lodging table')
    out={}
    for row in rows:
        if not row:continue
        name=prefecture(row[0])
        if not name:continue
        if name in out:raise ValueError('Duplicate lodging prefecture '+name)
        if len(row)<7:raise ValueError('Short lodging source row')
        out[name]=[number(str(x).strip().lstrip('*')) for x in row[1:7]]
    if len(out)!=48:raise ValueError('Expected national+47 lodging prefectures')
    return out


def lodging_tourism(raw,year):
    import openpyxl
    w=openpyxl.load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
    out={p:{} for p in PREFECTURES}
    for name,columns in [('第2表(年計)',{'total_stays':1,'foreign_stays':16}),
                         ('第8表(年計)',{'occupancy_rate':1,'ryokan_occupancy':16,'resort_hotel_occupancy':17,'business_hotel_occupancy':18,'city_hotel_occupancy':19})]:
        rows=list(w[name].values);seen=set()
        text=''.join(compact(v) for row in rows[:15] for v in row)
        # A year label is printed among the first data rows of each annual sheet.
        markers=[str(year), '令和元年' if year==2019 else '令和'+str(year-2018)+'年']
        if not any(m in text for m in markers):raise ValueError('Tourism observation year changed '+str(year))
        for row in rows:
            if not row:continue
            label=re.sub(r'^\d{2}','',compact(row[0]))
            if label not in PREFECTURES:continue
            if label in seen:raise ValueError('Duplicate lodging prefecture '+label)
            seen.add(label)
            for field,col in columns.items():out[label][field]=number(row[col])
        if seen!=set(PREFECTURES):raise ValueError('Incomplete tourism prefecture coverage')
    return out


def hotels(c,q):
    source_specs=json.loads(Path(__file__).with_name('hotel-sources.json').read_text())
    by_mhlw={}
    for s in source_specs:
        q.result['scopes'].append(dict(source=s['url'],year=s['year'],period='fiscal year',
            scope='47 prefectures: facilities, rooms, permits, closures; prior-year stock for closure rates'))
        try:by_mhlw[s['year']]=lodging_facilities(fetch(c,s['url']),s['year'])
        except Exception as e:by_mhlw[s['year']]={};q.error(s['url'],e)
    by_jta={}
    for year,file_id in [(2019,'002026792'),(2021,'002026777'),(2025,'002026762')]:
        url=f'https://www.mlit.go.jp/kankocho/content/{file_id}.xlsx'
        q.result['scopes'].append(dict(source=url,year=year,period='calendar year, annual final',
            scope='47 prefectures: total/foreign stays and five occupancy series',unit='person-nights; percent'))
        try:by_jta[year]=lodging_tourism(fetch(c,url),year)
        except Exception as e:by_jta[year]={};q.error(url,e)
    def diff(a,b):return a-b if a is not None and b is not None else None
    def ratio(a,b):return round(a/b,4) if a is not None and b not in [None,0] else None
    def percent(a,b):return round(a/b*100,2) if a is not None and b not in [None,0] else None
    def add(values):return sum(values) if all(x is not None for x in values) else None
    data=load(c,SITE+'/maps/lodging-trends/ja/data/prefecture_master.json')
    if len(data)!=47:raise ValueError('Published lodging prefecture count !=47')
    for row in data:
        name=PREFECTURES[int(row['prefecture_code'])-1];original={}
        for y in [2019,2021,2024]:
            v=by_mhlw.get(y,{}).get(name,[None]*6)
            for i,field in enumerate(['hotel_facilities','hotel_rooms','simple_lodging_facilities','boarding_house_facilities','permits_all_lodging','closures_all_lodging']):
                original[field+'_fy'+str(y)]=v[i]
            prior=y-1;pv=by_mhlw.get(prior,{}).get(name,[None]*6)
            stock=add([pv[i] for i in [0,2,3]])
            original['total_lodging_end_fy'+str(prior)]=stock
            original['closure_rate_all_lodging_fy'+str(y)]=ratio(v[5],stock)
            original['flow_diff_all_lodging_fy'+str(y)]=diff(v[4],v[5])
        for y in [2019,2021,2025]:
            s=by_jta.get(y,{}).get(name,{})
            for field in ['total_stays','foreign_stays','occupancy_rate','ryokan_occupancy','resort_hotel_occupancy','business_hotel_occupancy','city_hotel_occupancy']:
                original[field+'_'+str(y)]=s.get(field)
            jp=diff(s.get('total_stays'),s.get('foreign_stays'));original['japanese_stays_'+str(y)]=jp
            original['foreign_ratio_'+str(y)]=ratio(s.get('foreign_stays'),s.get('total_stays'))
            original['japanese_ratio_'+str(y)]=ratio(jp,s.get('total_stays'))
        for y in [2021,2024]:
            original['hotel_stock_change_'+str(y)+'_2019']=diff(original['hotel_facilities_fy'+str(y)],original['hotel_facilities_fy2019'])
            original['hotel_rooms_change_'+str(y)+'_2019']=diff(original['hotel_rooms_fy'+str(y)],original['hotel_rooms_fy2019'])
            original['hotel_facilities_index_fy'+str(y)]=percent(original['hotel_facilities_fy'+str(y)],original['hotel_facilities_fy2019'])
        original['hotel_rooms_index_fy2024']=percent(original['hotel_rooms_fy2024'],original['hotel_rooms_fy2019'])
        original['hotel_rooms_change_rate_fy2024']=ratio(original['hotel_rooms_change_2024_2019'],original['hotel_rooms_fy2019'])
        for y in [2021,2025]:
            for outfield,infield in [('foreign_recovery_rate','foreign_stays'),('stays_recovery_rate','total_stays'),('japanese_recovery_rate','japanese_stays')]:
                a,b=original[infield+'_'+str(y)],original[infield+'_2019']
                original[outfield+'_'+str(y)]=round(a/b*100,2) if a is not None and b not in [None,0] else None
            d=diff(original['occupancy_rate_'+str(y)],original['occupancy_rate_2019'])
            original['occupancy_pp_diff_'+str(y)+'_2019']=round(d,1) if d is not None else None
        for field,value in original.items():q.cell(name,field,row.get(field),value)
    q.result['limitations'] += ['Fiscal-year facility/permit/closure data and calendar-year stay/occupancy data are compared independently, not treated as the same period.',
        'Stays count person-nights, not unique travellers. Closure notifications are not necessarily bankruptcies.',
        'Annual tourism files can be revised after publication; mismatches require edition review before correction.',
        'Derived shares use four decimals, recovery indices use percent with two decimals, and occupancy differences use percentage points with one decimal.',
        'This checks the shared Japanese prefecture master, not regional summary cards or other language copies.']


EXTRA={'minerals':minerals,'medical':medical,'port':ports,'hotel_ryokan':hotels}
