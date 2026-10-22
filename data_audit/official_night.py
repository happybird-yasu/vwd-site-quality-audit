"""Independent official-source comparisons. GET only; never writes the website.

Each declared scope has a denominator. Missing/suppressed/ambiguous source cells
are never converted to zero or counted as matches. Raw HTTP evidence is retained.
"""
from __future__ import annotations
import argparse
from collections import defaultdict
import csv
import io
import json
from pathlib import Path
import re
import zipfile
import unicodedata
import logging

from .run import Client, SITE, now, write

PREFECTURES='北海道 青森県 岩手県 宮城県 秋田県 山形県 福島県 茨城県 栃木県 群馬県 埼玉県 千葉県 東京都 神奈川県 新潟県 富山県 石川県 福井県 山梨県 長野県 岐阜県 静岡県 愛知県 三重県 滋賀県 京都府 大阪府 兵庫県 奈良県 和歌山県 鳥取県 島根県 岡山県 広島県 山口県 徳島県 香川県 愛媛県 高知県 福岡県 佐賀県 長崎県 熊本県 大分県 宮崎県 鹿児島県 沖縄県'.split()


def number(v):
    if isinstance(v, bool) or v is None:
        return None
    try:
        s = str(v).strip().replace(',', '')
        return float(s) if re.fullmatch(r'-?\d+(?:\.\d+)?', s) else None
    except (ValueError, TypeError):
        return None


class Comparison:
    def __init__(self, theme):
        self.result = dict(theme=theme, started_at=now(), expected=0, compared=0,
                           matched=0, mismatches=0, unverified=0, details=[], errors=[],
                           scopes=[], limitations=[], full_theme_verification=False)

    def cell(self, key, field, actual, source, reason=None):
        r = self.result
        r['expected'] += 1
        a, b = number(actual), number(source)
        if reason or a is None or b is None:
            r['unverified'] += 1
            status = 'unverified'
            reason = reason or 'non_numeric_or_suppressed_value'
        else:
            r['compared'] += 1
            status = 'matched' if a == b else 'mismatch'
            r['matched' if a == b else 'mismatches'] += 1
        # All mismatches and missing cells, not a truncated list.
        if status != 'matched':
            r['details'].append(dict(key=key, field=field, published=actual,
                                     official=source, status=status, reason=reason))

    def error(self, scope, e):
        self.result['errors'].append(dict(scope=scope, error=type(e).__name__+': '+str(e)[:700]))

    def finish(self):
        r = self.result
        r['finished_at'] = now()
        r['status'] = ('issues_found' if r['mismatches'] else
                       'incomplete' if r['errors'] or r['unverified'] or not r['compared'] else
                       'declared_numeric_scope_matched')
        return r


def fetch(c, url, kind='official_numeric_source'):
    b, record = c.get(url, kind)
    if b is None:
        raise RuntimeError(record.get('error', 'fetch unavailable')+' '+url)
    return b


def load(c, url):
    return json.loads(fetch(c, url, 'published_numeric_data'))


def station_index(raw):
    z = zipfile.ZipFile(io.BytesIO(raw))
    names = [n for n in z.namelist() if '/UTF-8/' in n and n.endswith('.geojson')]
    if len(names) != 1:
        raise ValueError('Expected one UTF-8 official station GeoJSON')
    o = json.loads(z.read(names[0]))
    rows = defaultdict(list)
    for feature in o['features']:
        p = feature['properties']
        if 'S12_061' not in p:
            raise ValueError('2024 passenger field missing: wrong edition/schema')
        rows[(p['S12_001'], p['S12_002'], p['S12_003'])].append(p)
    return rows


def stations(c, q):
    source = 'https://nlftp.mlit.go.jp/ksj/gml/data/S12/S12-25/S12-25_GML.zip'
    q.result['scopes'].append(dict(source=source, years=[2019,2021,2023,2024],
        key='station name + operator + line', unit='passengers/day',
        data_availability_codes='S12 passenger availability retained in mismatch evidence'))
    index = station_index(fetch(c, source))
    for edition, filename, fields in [
        ('2024','all_stations_light_v2.json', {'f19':'041','f21':'049','f24':'061'}),
        ('2023','all_stations_light_v2_2023.json', {'f19':'041','f21':'049','f23':'057'})]:
        url = f'{SITE}/maps/station/{edition}/{filename}'
        try:
            data = load(c, url)
            if not isinstance(data, list) or not data:
                raise ValueError('No published station rows')
            for i, row in enumerate(data):
                key = (row['name'], row['operator'], row['line'])
                candidates = index.get(key, [])
                for field, suffix in fields.items():
                    if field not in row:
                        continue
                    # Duplicate official rail geometries can share a station key.
                    # Only compare when ALL candidates agree; never cherry-pick a match.
                    values = {p.get('S12_'+suffix) for p in candidates}
                    availability = {p.get('S12_'+str(int(suffix)-2).zfill(3)) for p in candidates}
                    reason = ('official_station_key_missing' if not candidates else
                              'ambiguous_official_station_key' if len(values)!=1 else
                              'official_no_passenger_data_code' if availability != {1} else None)
                    q.cell([edition,i,*key], field, row[field], next(iter(values)) if len(values)==1 else None, reason)
        except Exception as e:
            q.error(url,e)
    q.result['limitations'] += [
        'Station source editions can revise historical values: mismatches require edition review.',
        'Ranking cards, region assignment, coordinates and other language copies remain outside this numerical scope.',
        'Official absence codes are not zero passengers. These cells remain explicitly unverified.']


def station_ranking(c,q):
    source='https://nlftp.mlit.go.jp/ksj/gml/data/S12/S12-25/S12-25_GML.zip'
    index=station_index(fetch(c,source))
    q.result['scopes'].append(dict(source=source,years=[2019,2021,2023,2024],
        fields=['f19','f21','f23','f24'],unit='passengers/day',
        scope='all station passenger values in Japanese and multilingual ranking data sources'))
    for base in ['/cards/station-ja/assets/data/','/cards/station/assets/data/']:
        try:
            data=load(c,SITE+base+'stations.json')
            columns=data['meta']['fields']
            if len(data['rows'])!=data['meta']['rows']:raise ValueError('Ranking row count changed')
            for values in data['rows']:
                if len(values)!=len(columns):raise ValueError('Ranking field count changed')
                row=dict(zip(columns,values));key=(row['name'],row['operator'],row['line'])
                candidates=index.get(key,[])
                for field,suffix in [('f19','041'),('f21','049'),('f23','057'),('f24','061')]:
                    nums={p.get('S12_'+suffix) for p in candidates}
                    codes={p.get('S12_'+str(int(suffix)-2).zfill(3)) for p in candidates}
                    reason=('official_station_key_missing' if not candidates else
                            'ambiguous_official_station_key' if len(nums)!=1 else
                            'official_no_passenger_data_code' if codes!={1} else None)
                    q.cell([base,row['id'],*key],field,row.get(field),next(iter(nums)) if len(nums)==1 else None,reason)
        except Exception as e:q.error(base,e)
    q.result['limitations'] += ['Passenger values are checked against official records; rank order, geography, tie handling and HTML display remain unverified.',
        'Historical revisions in later official editions require source-edition review before any correction.',
        'No source availability code is interpreted as an observed zero.']


def road_index(raw):
    rows = list(csv.reader(io.StringIO(raw.decode('cp932'))))
    if not rows or len(rows[0]) != 159 or not any('令和３年度' in c for c in rows[0]):
        raise ValueError('Unexpected official road CSV schema/year')
    result = {}
    for rownum, row in enumerate(rows[1:],2):
        if len(row)!=159:
            raise ValueError(f'Road source row {rownum}: expected159 columns')
        key=row[0].strip().zfill(11)
        if key in result:
            raise ValueError('Duplicate road section '+key)
        result[key] = dict(weekday_24h_total_2021=row[61], weekday_24h_total_2015=row[69],
                           applied_section_length_km=row[23])
    return result


def roads(c,q):
    fields = ['weekday_24h_total_2021','weekday_24h_total_2015','applied_section_length_km']
    regions=['hokkaido','tohoku','kanto','koshinetsu','chubu-hokuriku','kinki','chugoku','shikoku','kyushu-okinawa']
    cache={}
    q.result['scopes'].append(dict(source='https://www.mlit.go.jp/road/census/r3/data/csv/kasyo{01..47}.csv',
        years=[2015,2021], fields=fields, key='traffic basic section ID',
        unit='vehicles/24h; km', published_scope='all 9 Japanese regional point files'))
    for region in regions:
        url=f'{SITE}/maps/road/data-20261004-speed1/{region}/points.json'
        try:
            data=load(c,url)
            if not data.get('features'):raise ValueError('No road features')
            for f in data['features']:
                p=f['properties'];key=str(p['basic_section_id']).zfill(11);code=key[:2]
                if code not in cache:
                    try:
                        cache[code]=road_index(fetch(c,f'https://www.mlit.go.jp/road/census/r3/data/csv/kasyo{code}.csv'))
                    except Exception as e:
                        cache[code]={};q.error('official prefecture '+code,e)
                original=cache[code].get(key)
                for field in fields:
                    q.cell([region,key],field,p.get(field),(original or {}).get(field),
                           None if original else 'official_section_missing_or_source_failed')
        except Exception as e:q.error(region,e)
    q.result['limitations'] += ['Geometry, counter positions, translations and derived ranking cards are not verified by matching traffic counts.',
        '2015 values are the previous-survey comparison recorded in the 2021 official table, not a separate 2015 boundary reconstruction.',
        'Source blank/suppressed values remain unverified, including published nulls.']


def onsen_rows(text):
    if '令和６年度' not in text or '令和７年３月末' not in text:
        raise ValueError('Unexpected onsen statistical year')
    result={}
    for line in text.splitlines():
        parts=line.split()
        if not parts or not re.fullmatch(r'北海道|\S+[都府県]',parts[0]):continue
        v=[number(x) for x in parts[1:]]
        if len(v)<15 or any(x is None for x in v[:15]):raise ValueError('Malformed prefecture row '+line)
        if parts[0] in result:raise ValueError('Duplicate prefecture')
        result[parts[0]]=dict(onsen_areas=v[2],total_springs=v[3],used_springs=v[4]+v[5],
                               unused_springs=v[6]+v[7],flow_l_min=v[12])
    if len(result)!=47:raise ValueError(f'Expected47 onsen prefectures, got{len(result)}')
    return result


def onsen(c,q):
    import pdfplumber
    logging.getLogger('pdfminer').setLevel(logging.ERROR)
    source='https://www.env.go.jp/nature/onsen/pdf/6-7_p_1.pdf'
    raw=fetch(c,source)
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        if len(pdf.pages)!=1:raise ValueError('Unexpected PDF page count')
        original=onsen_rows(pdf.pages[0].extract_text())
    q.result['scopes'].append(dict(source=source,year='FY2024, as of March2025',
        fields=list(next(iter(original.values()))),unit='areas; springs; L/min'))
    for lang in ['ja','en','de','es','fr']:
        url=f'{SITE}/maps/onsen/{lang}/national-data-20260918.json'
        try:
            data=load(c,url)['ENV']
            if len(data)!=47:raise ValueError('Published onsen prefecture count !=47')
            for code,row in data.items():
                src=original.get(PREFECTURES[int(code)-1]) if re.fullmatch(r'0[1-9]|[1-3][0-9]|4[0-7]',code) else None
                for field in next(iter(original.values())):
                    q.cell([lang,code,row['name']],field,row.get(field),(src or {}).get(field),None if src else 'prefecture_missing')
        except Exception as e:q.error(url,e)
    q.result['limitations'] += ['Historical GSJ chemical sampling points (POINTS), their temperature/pH/date/coordinates and regional cards remain unverified.',
        'Toyama original table has total169 but used126 + unused44 =170; comparison preserves source discrepancy without altering either value.']


def compact(v):
    return re.sub(r'\s+','',unicodedata.normalize('NFKC',str(v or '')))


def airport_rows(raw):
    import openpyxl
    w=openpyxl.load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
    ws=w['暦年・年度別空港管理状況調書'];rows=list(ws.values);out={}
    # Each airport begins with name and passenger column headings. Only the first
    # ten CALENDAR rows count: following fiscal-year/cargo rows are excluded.
    for i,row in enumerate(rows):
        if len(row)<13 or compact(row[1])!='空港名:':continue
        if 'ヘリ' in compact(row[2]) or 'へリ' in compact(row[2]):continue
        name=re.split('[（(]',str(row[2]))[0].strip(); block=rows[i:i+16]
        header=''.join(compact(v) for r in block[:4] for v in r)
        if '乗降客数' not in header or '国際線' not in header:raise ValueError('Airport passenger schema changed')
        for offset,year in [(9,2019),(11,2021),(13,2023),(14,2024)]:
            r=rows[i+offset];label=compact(r[1]);expected='元' if year==2019 else str(year-2018)
            if label!=expected:raise ValueError(f'Airport calendar row changed {name} {year}: {label}')
            key=(compact(name),year)
            if key in out:raise ValueError('Duplicate airport calendar record '+str(key))
            out[key]=dict(intl=r[8],dom=r[11],total=r[12])
    if len(out)<300:raise ValueError('Too few airport rows')
    return out


def airports(c,q):
    source='https://www.mlit.go.jp/koku/content/001905536.xlsx'
    original=airport_rows(fetch(c,source))
    q.result['scopes'].append(dict(source=source,years=[2019,2021,2023,2024],
        fields=['intl','dom','total'],unit='passengers/calendar year',
        definition='International subtotal includes transit passengers as in official source.'))
    data=load(c,SITE+'/maps/airport/airports.json')
    for row in data:
        for year in [2019,2021,2023,2024]:
            src=original.get((compact(row['name']),year))
            for field in ['intl','dom','total']:
                q.cell([row['name'],year],field,(row.get('p'+str(year)) or {}).get(field),(src or {}).get(field),
                       None if src else 'official_airport_name_not_matched')
    q.result['limitations'] += ['Coordinates, article text and derived cards are outside this numerical comparison.',
        'Blank source cells remain unverified; never inferred as zero.']


def workbook_rows(raw):
    if raw.startswith(b'PK'):
        import openpyxl
        w=openpyxl.load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
        return [list(s.values) for s in w]
    import xlrd
    w=xlrd.open_workbook(file_contents=raw)
    return [[s.row_values(i) for i in range(s.nrows)] for s in w.sheets()]


def school_table(raw,year,kind):
    sheets=workbook_rows(raw);high=kind=='高等学校'
    # Column locations refer to the original official tables, not exported CSVs.
    if high:
        label,total,pub,private=(0,1,7,11) if year==1965 else (1,2,7,31)
        main_label,main_col,branch_label,branch_col=(14,0,0,1) if year==1965 else (1,2,1,2)
        nat={1965:6,1990:8,2024:7}[year]
    else:
        label,total,pub,private=(0,1,5,21) if year==1965 else (1,2,6,24 if year==2024 else 21)
        main_col,branch_col=(2,3) if year==1965 else (3,4)
        nat=({1965:9,1990:8,2024:7} if kind=='小学校' else {1965:7,1990:8,2024:7})[year]
    national_label=compact(sheets[0][nat][label])
    expected={1965:'昭和40年度',1990:'平成2年度',2024:'計'}[year]
    if not national_label.startswith(expected):raise ValueError('School source year/total row changed '+national_label)
    if year==2024 and compact(sheets[0][nat][0])!='令和6年度':raise ValueError('School year !=2024')
    aliases={compact(p):p for p in PREFECTURES};aliases.update({compact(p[:-1]):p for p in PREFECTURES if p!='北海道'})
    def prefkey(v):return aliases.get(re.sub(r'^\d+','',compact(v)))
    def bypref(sheet,col):
        out={}
        for row in sheet:
            if len(row)<=col:continue
            pref=prefkey(row[col])
            if pref:
                if pref in out:raise ValueError('Duplicate school prefecture '+pref)
                out[pref]=row
        wanted=set(PREFECTURES[:46 if year==1965 else 47])
        if set(out)!=wanted:raise ValueError('School prefecture coverage changed')
        return out
    totals=bypref(sheets[0],label)
    mains=bypref(sheets[1],main_label) if high else totals
    branches=bypref(sheets[2],branch_label) if high else totals
    out={}
    for pref,row in totals.items():
        t,p,v=number(row[total]),number(row[pub]),number(row[private])
        out[pref]=dict(school_count_total=row[total],public_school_count=row[pub],private_school_count=row[private],
            national_school_count=t-p-v if all(x is not None for x in [t,p,v]) else None,
            main_school_count=mains[pref][main_col],branch_school_count=branches[pref][branch_col])
    return out


def schools(c,q):
    from .checks import json_values
    raw=fetch(c,SITE+'/maps/schools/assets/data.js','published_numeric_data')
    objects=json_values(raw.decode('utf-8'),is_json=False)[0]
    data=next(o for name,o in objects if isinstance(o,dict) and 'prefectureRows' in o)
    sources=json.loads(Path(__file__).with_name('school-sources.json').read_text())
    lookup={}
    for s in sources:
        key=(s['year'],s['type'])
        if key in lookup:raise ValueError('Duplicate school source specification')
        q.result['scopes'].append(dict(source=s['url'],year=s['year'],school_type=s['type'],unit='schools',
            scope='prefecture school counts: total, main, branch, national, public, private'))
        try:lookup[key]=school_table(fetch(c,s['url']),*key)
        except Exception as e:lookup[key]={};q.error(str(key),e)
    for row in data['prefectureRows']:
        kind='高等学校' if row['school_type'].startswith('高等学校') else row['school_type']
        src=lookup.get((row['year'],kind),{}).get(row['prefecture'])
        for field in ['school_count_total','main_school_count','branch_school_count','national_school_count','public_school_count','private_school_count']:
            reason=None if src else ('1965_Okinawa_outside_survey_scope' if row['year']==1965 and row['prefecture']=='沖縄県' else 'official_prefecture_missing_or_source_failed')
            if row['year']==1965 and kind=='小学校' and row['prefecture']=='神奈川県' and field in ['private_school_count','national_school_count']:
                reason='Published documented private-school correction51to31; retain raw-source discrepancy for review'
            q.cell([row['year'],kind,row['prefecture']],field,row.get(field),(src or {}).get(field),reason)
    q.result['limitations'] += ['National student counts, national summary and per-school averages are not covered by these prefecture school-count source tables.',
        '1965 excludes Okinawa; absent coverage is explicitly unverified, never treated as zero.',
        'The documented Kanagawa1965 private-school correction is flagged separately, not silently accepted as a source match.',
        'Derived national ownership counts depend on original total/public/private; blank values remain unverified.']


ADAPTERS={'station':stations,'station_ranking':station_ranking,'road':roads,'onsen':onsen,'airport':airports,'schools':schools}


def main():
    p=argparse.ArgumentParser();p.add_argument('--theme',required=True);p.add_argument('--out',type=Path,required=True)
    args=p.parse_args();q=Comparison(args.theme);c=Client(args.out,delay=3,minutes=300)
    try:
        from .official_extra import EXTRA
        adapter=ADAPTERS.get(args.theme) or EXTRA.get(args.theme)
        if adapter:adapter(c,q)
        else:q.result['limitations'].append('Official numerical adapter not yet implemented. This theme is UNVERIFIED, not passed.')
    except Exception as e:q.error(args.theme,e)
    finally:
        r=q.finish();write(args.out/'official-report.json',r);write(args.out/'fetch_manifest.json',c.records)
        write(args.out/'robots_diagnostics.json',{'sources':c.robot_sources,'errors':c.robot_errors})
        lines=[f"## {args.theme}: {r['status']}",f"Expected {r['expected']}; compared {r['compared']}; matched {r['matched']}; mismatches {r['mismatches']}; unverified {r['unverified']}; errors {len(r['errors'])}.",
               'Full theme verification: **not complete** (see scope and limitations).',*['- '+x for x in r['limitations']]]
        (args.out/'official-report.md').write_text('\n\n'.join(lines),encoding='utf-8')
        print(json.dumps({k:v for k,v in r.items() if k not in ['details']},ensure_ascii=False),flush=True)


if __name__=='__main__':main()
