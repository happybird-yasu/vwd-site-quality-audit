import io
import json
import unittest
import zipfile
from .official_night import Comparison, number, road_index, station_index, onsen_rows, school_table


class OfficialTests(unittest.TestCase):
    def test_missing_and_suppressed_are_not_matches(self):
        q=Comparison('test')
        q.cell('a','count',0,'-')
        q.cell('b','count',None,None)
        q.cell('c','count',0,0,'official_no_passenger_data_code')
        q.cell('d','count',17,18)
        q.cell('e','count',1234,'1,234')
        r=q.finish()
        self.assertEqual((r['expected'],r['compared'],r['mismatches'],r['unverified']),(5,2,1,3))
        self.assertFalse(r['full_theme_verification'])

    def test_zero_work_never_passes(self):
        self.assertEqual(Comparison('empty').finish()['status'],'incomplete')
        for v in ['N/A','…','-',float('nan'),True]:self.assertIsNone(number(v))

    def test_official_leading_decimal_is_numeric(self):
        self.assertEqual(number('.1'),0.1)
        self.assertEqual(number('-.8'),-0.8)
        self.assertIsNone(number('.'))

    def test_covid_official_japanese_name_alias(self):
        from .infection import csv_table
        b='2021年23週\n,新型コロナウイルス感染症,,Ｅ型肝炎,\n,報告,累積,報告,累積\n総数,5,575,9,100\n'.encode('cp932')
        self.assertEqual(csv_table(b,'COVID-19',2021,23)['総数']['count'],'5')
        self.assertEqual(csv_table(b,'Ｅ型肝炎',2021,23)['総数']['count'],'9')

    def test_aggregate_rate_mean_is_not_sum(self):
        from .infection import compare_aggregates
        class Fixture:
            def get(self,url,kind):
                week=int(url.split('-')[-2]);rate={23:0.1,24:0.2,25:0.2}[week]
                return (f'2021年{week}週\n,病名,\n,報告,定当\n総数,2,{rate}\n'.encode('cp932'),{})
        data={'periodType':'month','weeks':[23,24,25],'sources':[f'https://example.go.jp/2021-{w}-teiten.csv' for w in [23,24,25]],'prefs':{'00':{'rate':0.1667,'count':6}}}
        catalog={'seriesIndex':{'2021-m06-病名':{'path':'data/m06.json'}},'prefs':{}}
        r=compare_aggregates(Fixture(),catalog,[('https://example.com/data/m06.json',data)])
        self.assertEqual((r['numeric_values_compared'],r['numeric_mismatches']),(2,0))

    def test_road_year_and_duplicate_guard(self):
        import csv
        header=['']*159;header[0]='令和３年度';row=['']*159
        row[0]='08111200010';row[61]='12034';row[69]='12468';row[23]='2.2'
        def data(rows):
            s=io.StringIO();csv.writer(s).writerows(rows);return s.getvalue().encode('cp932')
        self.assertEqual(road_index(data([header,row]))[row[0]]['weekday_24h_total_2015'],'12468')
        with self.assertRaises(ValueError):road_index(data([header,row,row]))
        with self.assertRaises(ValueError):road_index(data([['平成27年度']*159,row]))

    def test_station_wrong_edition_and_ambiguity_retained(self):
        def data(props):
            b=io.BytesIO()
            with zipfile.ZipFile(b,'w') as z:z.writestr('x/UTF-8/x.geojson',json.dumps({'features':[{'properties':p} for p in props]}))
            return b.getvalue()
        p={'S12_001':'同名駅','S12_002':'会社A','S12_003':'路線A','S12_061':12}
        index=station_index(data([p,dict(p,S12_061=13)]))
        self.assertEqual(len(index[('同名駅','会社A','路線A')]),2)
        del p['S12_061']
        with self.assertRaises(ValueError):station_index(data([p]))

    def test_onsen_source_year_and_prefecture_count(self):
        with self.assertRaises(ValueError):onsen_rows('令和５年度')
        with self.assertRaises(ValueError):onsen_rows('令和６年度 令和７年３月末\n富山県 9 15 61 169 52 74 13 31 36 64 66 0 29181 17800 11381')

    def test_school_wrong_year_is_rejected(self):
        import openpyxl
        w=openpyxl.Workbook();s=w.active;s.cell(8,2,'計');s.cell(8,1,'令和5年度')
        b=io.BytesIO();w.save(b)
        with self.assertRaisesRegex(ValueError,'2024'):school_table(b.getvalue(),2024,'小学校')


if __name__=='__main__':unittest.main()
