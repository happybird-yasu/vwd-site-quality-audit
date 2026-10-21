"""Official source handling: e-Stat reference tables of the Report on Internal
Migration (住民基本台帳人口移動報告 参考表).

kind "nationality": 第3表(転入)/第4表(転出)  総数・日本人・外国人
kind "agesex":      第1表(転入)/第2表(転出)  年齢(10歳階級)・男女
"""
from __future__ import annotations

import re
import urllib.parse
from dataclasses import dataclass, field
from pathlib import Path

from .common import http_get, sha256_bytes, sha256_file, utcnow

LIST_URL = (
    "https://www.e-stat.go.jp/stat-search/files?page=1&layout=datalist"
    "&toukei=00200523&tstat=000000070001&tclass1={tclass}&cycle=7"
    "&year={year}0&month=0&tclass2val=0"
)
DOWNLOAD_URL = "https://www.e-stat.go.jp/stat-search/file-download?statInfId={sid}&fileKind={fk}"
TCLASS = {"nationality": "000001128855", "agesex": "000001128355"}
ESTAT_INTERVAL = 2.0

PREF_NAMES = [
    "北海道", "青森県", "岩手県", "宮城県", "秋田県", "山形県", "福島県", "茨城県", "栃木県",
    "群馬県", "埼玉県", "千葉県", "東京都", "神奈川県", "新潟県", "富山県", "石川県", "福井県",
    "山梨県", "長野県", "岐阜県", "静岡県", "愛知県", "三重県", "滋賀県", "京都府", "大阪府",
    "兵庫県", "奈良県", "和歌山県", "鳥取県", "島根県", "岡山県", "広島県", "山口県", "徳島県",
    "香川県", "愛媛県", "高知県", "福岡県", "佐賀県", "長崎県", "熊本県", "大分県", "宮崎県",
    "鹿児島県", "沖縄県",
]

# Aggregate (non-municipality) origin rows: prefecture subtotals end in 000;
# designated-city totals and the 23-ward total are listed explicitly.
DESIGNATED_CITY_CODES = {
    "01100", "04100", "11100", "12100", "13100", "14100", "14130", "14150", "15100",
    "22100", "22130", "23100", "26100", "27100", "27140", "28100", "33100", "34100",
    "40100", "40130", "43100",
}

# Column layout (0-based). Verified by identity checks on every sheet.
COL_CODE = 8
COL_ALL = 9
NAT_COLS = {"japanese": 10, "foreign": 11}
AGE_ALL_COLS = list(range(10, 18))        # 0-9,10-19,20-29,30-39,40-49,50-59,60+,unknown
MALE_TOTAL, FEMALE_TOTAL = 18, 27
MALE_AGE_COLS = list(range(19, 27))
FEMALE_AGE_COLS = list(range(28, 36))
# public-site buckets -> official age columns (all-sexes block)
PUBLIC_AGE_BUCKETS = {"0-19": (10, 11), "20-39": (12, 13), "40-59": (14, 15), "60+": (16,)}
UNKNOWN_AGE_COL = 17


class LayoutError(Exception):
    pass


def is_aggregate_code(code: str) -> bool:
    return code.endswith("000") or code in DESIGNATED_CITY_CODES


@dataclass
class Listing:
    year: int
    kind: str
    url: str
    html_sha256: str
    items: list = field(default_factory=list)   # dicts
    warnings: list = field(default_factory=list)


def discover(year: int, kind: str) -> Listing:
    url = LIST_URL.format(tclass=TCLASS[kind], year=year)
    body, _, _ = http_get(url, host_key="estat", min_interval=ESTAT_INTERVAL)
    text = body.decode("utf-8", "ignore")
    found = re.findall(
        r'js-data"[^>]*>\s*([^<]+?)\s*</a>.*?file-download\?statInfId=(\d+)&(?:amp;)?fileKind=(\d)',
        text, re.S)
    lst = Listing(year=year, kind=kind, url=url, html_sha256=sha256_bytes(body))
    if len(found) != 94:
        lst.warnings.append(f"expected 94 list entries, found {len(found)}")
    for i, (name, sid, fk) in enumerate(found):
        direction = "inflow" if i < 47 else "outflow"
        pref_index = i % 47
        expected = PREF_NAMES[pref_index] if pref_index < len(PREF_NAMES) else None
        lst.items.append({
            "year": year, "kind": kind, "list_index": i, "direction_assumed": direction,
            "pref_code": f"{pref_index + 1:02d}", "pref_name_listed": name.strip(),
            "pref_name_expected": expected, "stat_inf_id": sid, "file_kind": fk,
            "url": DOWNLOAD_URL.format(sid=sid, fk=fk),
            "order_ok": name.strip() == expected,
        })
    if any(not it["order_ok"] for it in lst.items):
        lst.warnings.append("list order differs from JIS prefecture order for some entries")
    return lst


def download(item: dict, cache_dir: Path) -> dict:
    """Download one official file into cache (atomic), verify, return manifest row.

    Existing verified cache files are reused (not re-downloaded).
    """
    cache_dir.mkdir(parents=True, exist_ok=True)
    stem = f"{item['year']}_{item['kind']}_{item['stat_inf_id']}"
    for existing in cache_dir.glob(stem + ".*"):
        if existing.suffix in (".xls", ".xlsx") and existing.stat().st_size > 10_000:
            return {**item, "path": str(existing), "bytes": existing.stat().st_size,
                    "sha256": sha256_file(existing), "cached": True}
    body, headers, attempts = http_get(item["url"], host_key="estat", min_interval=ESTAT_INTERVAL)
    if body[:2] == b"PK":
        ext = ".xlsx"
    elif body[:4] == b"\xd0\xcf\x11\xe0":
        ext = ".xls"
    else:
        raise LayoutError(f"download is not an Excel file (got {len(body)} bytes, head={body[:8]!r})")
    if len(body) < 10_000:
        raise LayoutError(f"download too small: {len(body)} bytes")
    tmp = cache_dir / (stem + ext + ".part")
    tmp.write_bytes(body)
    final = cache_dir / (stem + ext)
    tmp.replace(final)
    disp = headers.get("content-disposition") or headers.get("Content-Disposition") or ""
    m = re.search(r"filename\*=UTF-8''([^;]+)", disp)
    fname = urllib.parse.unquote(m.group(1)) if m else None
    return {**item, "path": str(final), "bytes": len(body), "sha256": sha256_bytes(body),
            "served_filename": fname, "retrieved_at": utcnow(), "attempts": attempts,
            "cached": False}


# ---------------------------------------------------------------- parsing

def _num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else None
    return None


def col_letter(idx0: int) -> str:
    s, n = "", idx0 + 1
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


class Book:
    """Uniform access to .xls (xlrd) and .xlsx (openpyxl read-only)."""

    def __init__(self, path: Path):
        self.path = Path(path)
        self.is_xls = self.path.suffix.lower() == ".xls"
        if self.is_xls:
            import xlrd
            self._b = xlrd.open_workbook(str(path), on_demand=True)
            self.sheetnames = list(self._b.sheet_names())
        else:
            import openpyxl
            self._b = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            self.sheetnames = list(self._b.sheetnames)

    def rows(self, name: str):
        if self.is_xls:
            sh = self._b.sheet_by_name(name)
            for i in range(sh.nrows):
                yield tuple(sh.row_values(i))
        else:
            for r in self._b[name].iter_rows(values_only=True):
                yield tuple(r)

    def close(self):
        try:
            if self.is_xls:
                self._b.release_resources()
            else:
                self._b.close()
        except Exception:
            pass


@dataclass
class Sheet:
    name: str
    dest_code: str
    kind: str
    direction: str | None = None
    title: str = ""
    # origin_code -> {"row": 1-based, "all": int|None, ...}
    records: dict = field(default_factory=dict)
    non_numeric: list = field(default_factory=list)
    total_row: dict | None = None
    integrity: list = field(default_factory=list)   # failed official-internal identities
    sum_all_numeric: int = 0


def parse_sheet(name: str, rows, kind: str) -> Sheet:
    m = re.match(r"(\d{5})", name)
    if not m:
        raise LayoutError(f"sheet name has no 5-digit code: {name!r}")
    sh = Sheet(name=name, dest_code=m.group(1), kind=kind)
    ncols = 12 if kind == "nationality" else 36
    for idx, row in enumerate(rows, 1):
        if idx <= 14:
            for c in row:
                if isinstance(c, str):
                    if "転入者数" in c:
                        sh.direction = sh.direction or "inflow"
                        sh.title = sh.title or c.strip()
                    elif "転出者数" in c:
                        sh.direction = sh.direction or "outflow"
                        sh.title = sh.title or c.strip()
        if len(row) <= COL_ALL:
            continue
        code = row[COL_CODE]
        code = str(code).strip() if code is not None else ""
        if not re.fullmatch(r"\d{5}", code):
            continue
        row = tuple(row) + (None,) * max(0, ncols - len(row))
        vals = {}
        cols = [COL_ALL] + (list(NAT_COLS.values()) if kind == "nationality"
                             else AGE_ALL_COLS + [MALE_TOTAL] + MALE_AGE_COLS
                             + [FEMALE_TOTAL] + FEMALE_AGE_COLS)
        for c in cols:
            n = _num(row[c])
            vals[c] = n
            if n is None and row[c] not in (None, ""):
                sh.non_numeric.append({"row": idx, "col": col_letter(c), "raw": str(row[c])[:20]})
        label = re.sub(r"\s+", "", str(row[7])) if len(row) > 7 and row[7] is not None else ""
        rec = {"row": idx, "vals": vals, "label": label}
        if code == "00000":
            sh.total_row = rec
        sh.records[code] = rec
        if vals.get(COL_ALL) is not None:
            sh.sum_all_numeric += vals[COL_ALL]
    if sh.total_row is None:
        raise LayoutError(f"{name}: total row (area code 00000) not found")
    _identities(sh)
    return sh


def _identities(sh: Sheet) -> None:
    t = sh.total_row["vals"]
    ok = lambda *xs: all(x is not None for x in xs)  # noqa: E731
    if sh.kind == "nationality":
        a, j, f = t[COL_ALL], t[10], t[11]
        if ok(a, j, f) and a != j + f:
            sh.integrity.append(f"total {a} != japanese {j} + foreign {f}")
    else:
        a, m, f = t[COL_ALL], t[MALE_TOTAL], t[FEMALE_TOTAL]
        if ok(a, m, f) and a != m + f:
            sh.integrity.append(f"total {a} != male {m} + female {f}")
        ages = [t[c] for c in AGE_ALL_COLS]
        if ok(a, *ages) and sum(ages) != a:
            sh.integrity.append(f"age buckets sum {sum(ages)} != total {a}")
        mages = [t[c] for c in MALE_AGE_COLS]
        if ok(m, *mages) and sum(mages) != m:
            sh.integrity.append(f"male age sum {sum(mages)} != male total {m}")
        fages = [t[c] for c in FEMALE_AGE_COLS]
        if ok(f, *fages) and sum(fages) != f:
            sh.integrity.append(f"female age sum {sum(fages)} != female total {f}")
